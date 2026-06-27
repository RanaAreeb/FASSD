from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "fyp.txt"
OUT = ROOT / "FASSD_Final_FYP_Simple_Black_White.pptx"

WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(80, 80, 80)

ACRONYM = "Forensic Acoustics for Synthetic Speech Detection"


def add_textbox(slide, x, y, w, h, text="", size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_lines(slide, lines, x, y, w, h, size=18):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    first = True
    for raw in lines:
        line = clean_line(raw)
        if not line:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False

        if line.startswith("* "):
            p.text = line[2:]
            p.level = 0
            p.font.size = Pt(size)
        elif re.match(r"^\d+\.\s+", line):
            p.text = line
            p.level = 0
            p.font.size = Pt(size)
        elif line.startswith("|"):
            p.text = line
            p.font.name = "Consolas"
            p.font.size = Pt(max(size - 5, 9))
        elif line.startswith("## "):
            p.text = line.replace("## ", "")
            p.font.size = Pt(size + 3)
            p.font.bold = True
        elif line.startswith("### "):
            p.text = line.replace("### ", "")
            p.font.size = Pt(size + 1)
            p.font.bold = True
        else:
            p.text = line
            p.font.size = Pt(size)

        p.font.name = p.font.name or "Arial"
        p.font.color.rgb = BLACK
        p.space_after = Pt(5)

    return box


def clean_line(line: str) -> str:
    line = line.rstrip()
    line = line.replace("Forensic Audio Source and Spoofing Detector", ACRONYM)
    line = line.replace("**", "")
    return line


def parse_slides(text: str):
    parts = re.split(r"(?=## Slide \d+)", text)
    slides = []
    for part in parts:
        part = part.strip()
        if not part.startswith("## Slide "):
            continue

        lines = part.splitlines()
        header = clean_line(lines[0])
        m = re.match(r"## Slide \d+\s+[—-]\s+(.+)", header)
        title = m.group(1).strip() if m else header.replace("##", "").strip()

        body = []
        for line in lines[1:]:
            if line.startswith("**Speaker notes:**"):
                break
            if line.strip() == "---":
                break
            body.append(line)

        slides.append((title, body))
    return slides


def make_simple_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    source = SOURCE.read_text(encoding="utf-8")
    slides = parse_slides(source)

    blank = prs.slide_layouts[6]

    for idx, (slide_title, body) in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE

        if idx == 1:
            add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11.7), Inches(0.8), "FASSD", 44, True, BLACK, PP_ALIGN.CENTER)
            add_textbox(slide, Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.6), ACRONYM, 26, True, BLACK, PP_ALIGN.CENTER)
            add_textbox(slide, Inches(0.8), Inches(3.0), Inches(11.7), Inches(0.4), "Final Year Project Presentation", 20, False, BLACK, PP_ALIGN.CENTER)
            add_textbox(slide, Inches(0.8), Inches(4.1), Inches(11.7), Inches(0.8), "Team Members:\nRana M. Areeb\nM. Hasnain Siddique", 16, False, BLACK, PP_ALIGN.CENTER)
            add_textbox(slide, Inches(0.8), Inches(5.35), Inches(11.7), Inches(0.35), "Supervisor: Sir Faran Mehmood", 16, False, BLACK, PP_ALIGN.CENTER)
            add_textbox(slide, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.35), "https://www.deepfakedetection.dev/", 14, False, GRAY, PP_ALIGN.CENTER)
        else:
            add_textbox(slide, Inches(0.55), Inches(0.35), Inches(11.7), Inches(0.55), slide_title, 28, True, BLACK)
            visible_body = [line for line in body if line.strip()]
            add_lines(slide, visible_body, Inches(0.75), Inches(1.15), Inches(11.9), Inches(5.7), 16)

        add_textbox(slide, Inches(11.9), Inches(7.05), Inches(0.9), Inches(0.25), str(idx), 10, False, GRAY, PP_ALIGN.RIGHT)

    prs.save(OUT)
    print(f"Saved {OUT}")
    print(f"Slides: {len(slides)}")


if __name__ == "__main__":
    make_simple_deck()
