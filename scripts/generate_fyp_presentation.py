from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "FASSD_Final_FYP_Award_Style_Presentation_PowerPoint_Safe.pptx"

W, H = Inches(13.333), Inches(7.5)

BG = RGBColor(7, 12, 24)
PANEL = RGBColor(15, 23, 42)
PANEL_2 = RGBColor(20, 31, 52)
TEXT = RGBColor(241, 245, 249)
MUTED = RGBColor(148, 163, 184)
CYAN = RGBColor(34, 211, 238)
BLUE = RGBColor(59, 130, 246)
GREEN = RGBColor(52, 211, 153)
AMBER = RGBColor(251, 191, 36)
RED = RGBColor(248, 113, 113)
PURPLE = RGBColor(168, 85, 247)


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    # subtle accent orbs
    for x, y, color, alpha in [
        (Inches(10.8), Inches(-0.4), CYAN, 0.18),
        (Inches(-0.8), Inches(5.5), PURPLE, 0.15),
    ]:
        orb = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, Inches(2.8), Inches(2.8))
        orb.fill.solid()
        orb.fill.fore_color.rgb = color
        orb.fill.transparency = int((1 - alpha) * 100)
        orb.line.fill.background()


def textbox(slide, x, y, w, h, text="", size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def title(slide, text, kicker="FASSD Final Year Project", subtitle=None):
    textbox(slide, Inches(0.65), Inches(0.35), Inches(7.8), Inches(0.35), kicker.upper(), 9, True, CYAN)
    textbox(slide, Inches(0.65), Inches(0.72), Inches(10.5), Inches(0.72), text, 28, True, TEXT)
    if subtitle:
        textbox(slide, Inches(0.68), Inches(1.36), Inches(9.6), Inches(0.45), subtitle, 13, False, MUTED)


def footer(slide, n):
    textbox(slide, Inches(0.65), Inches(7.08), Inches(6), Inches(0.2), "Forensic Audio Source and Spoofing Detector", 8, False, MUTED)
    textbox(slide, Inches(12.1), Inches(7.08), Inches(0.6), Inches(0.2), f"{n:02d}", 8, True, CYAN, PP_ALIGN.RIGHT)


def bullet_list(slide, items, x, y, w, h, size=16, color=TEXT, gap=0.09):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(gap * 72)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def card(slide, x, y, w, h, heading, body, accent=CYAN, body_size=15):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL
    shp.line.color.rgb = accent
    shp.line.transparency = 35
    textbox(slide, x + Inches(0.22), y + Inches(0.16), w - Inches(0.44), Inches(0.3), heading.upper(), 9, True, accent)
    textbox(slide, x + Inches(0.22), y + Inches(0.52), w - Inches(0.44), h - Inches(0.65), body, body_size, True, TEXT)
    return shp


def quote(slide, text, x, y, w, h, accent=CYAN):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL_2
    shp.line.color.rgb = accent
    shp.line.width = Pt(1.5)
    textbox(slide, x + Inches(0.35), y + Inches(0.28), w - Inches(0.7), h - Inches(0.5), text, 23, True, TEXT, PP_ALIGN.CENTER)


def add_bar_chart(slide, x, y, w, h, title_text, categories, series, value_axis_title="", max_scale=None):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title_text
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = TEXT
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    if max_scale is not None:
        chart.value_axis.maximum_scale = max_scale
    if value_axis_title:
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = value_axis_title
    for ser in chart.series:
        ser.has_data_labels = True
        ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        ser.data_labels.number_format = '0.0'
    return chart


def add_doughnut_chart(slide, x, y, w, h, title_text, labels, values):
    data = ChartData()
    data.categories = labels
    data.add_series("Samples", values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, w, h, data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title_text
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = TEXT
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.show_percentage = True
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.BEST_FIT
    return chart


def add_table(slide, x, y, w, h, rows, cols, data, header_color=CYAN):
    table = slide.shapes.add_table(rows, cols, x, y, w, h).table
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL_2 if r == 0 else PANEL
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(10 if rows > 5 else 11)
                    run.font.bold = r == 0
                    run.font.color.rgb = header_color if r == 0 else TEXT
    return table


def flow(slide, labels, x, y, w, h, accent=CYAN):
    gap = Inches(0.12)
    box_w = (w - gap * (len(labels) - 1)) / len(labels)
    for i, lab in enumerate(labels):
        bx = x + i * (box_w + gap)
        card(slide, bx, y, box_w, h, f"{i + 1:02d}", lab, accent, 13)
        if i < len(labels) - 1:
            textbox(slide, bx + box_w - Inches(0.02), y + h / 2 - Inches(0.12), Inches(0.24), Inches(0.24), "→", 16, True, accent, PP_ALIGN.CENTER)


def notes(slide, text):
    try:
        ns = slide.notes_slide
        ns.notes_text_frame.text = text
    except Exception:
        pass


def new_slide(prs, n, t, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title(slide, t, subtitle=subtitle)
    footer(slide, n)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    n = 1

    # 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    textbox(s, Inches(0.7), Inches(0.55), Inches(2.2), Inches(0.35), "FINAL YEAR PROJECT", 11, True, CYAN)
    textbox(s, Inches(0.7), Inches(1.2), Inches(8.8), Inches(1.0), "FASSD", 58, True, TEXT)
    textbox(s, Inches(0.75), Inches(2.2), Inches(9.8), Inches(0.6), "Forensic Audio Source and Spoofing Detector", 24, True, CYAN)
    textbox(s, Inches(0.78), Inches(3.0), Inches(8.2), Inches(0.6), "Multi-axis forensic audio evidence reporting for AI-origin, replay, channel, and partial fabrication review.", 17, False, MUTED)
    card(s, Inches(0.75), Inches(4.35), Inches(3.15), Inches(0.95), "Team", "Rana M. Areeb\nM. Hasnain Siddique", CYAN, 13)
    card(s, Inches(4.1), Inches(4.35), Inches(2.7), Inches(0.95), "Supervisor", "Sir Faran Mehmood", GREEN, 14)
    card(s, Inches(7.0), Inches(4.35), Inches(4.8), Inches(0.95), "Live System", "https://www.deepfakedetection.dev", AMBER, 13)
    footer(s, n)
    notes(s, "Good morning respected panel. Our project is FASSD, an experimental forensic audio decision-support system.")
    n += 1

    # 2
    s = new_slide(prs, n, "Opening Story", "Why the project changed from a detector into a forensic evidence system")
    card(s, Inches(0.8), Inches(2.0), Inches(4.8), Inches(1.25), "Started as", "Detect AI-generated speech using ML", CYAN, 22)
    card(s, Inches(7.6), Inches(2.0), Inches(4.8), Inches(1.25), "Became", "Explain what evidence exists in audio", GREEN, 22)
    textbox(s, Inches(5.95), Inches(2.42), Inches(1.2), Inches(0.5), "→", 34, True, AMBER, PP_ALIGN.CENTER)
    bullet_list(s, ["Real audio may be compressed, replayed, mixed, edited, or re-uploaded.", "A single fake/real score cannot explain these different failure modes."], Inches(1.1), Inches(4.0), Inches(10.8), Inches(1.2), 18)
    quote(s, "Main realization: a single fake/real detector is not enough for forensic audio review.", Inches(1.15), Inches(5.45), Inches(11.0), Inches(0.9), AMBER)
    notes(s, "We started with binary deepfake detection. Testing showed real-world audio requires evidence separation.")
    n += 1

    # 3
    s = new_slide(prs, n, "Problem Statement", "Modern voice cloning is realistic; real recordings are messy")
    bullet_list(s, ["Noise", "Compression", "Replay", "Broadcast chains", "Speaker-to-device recording", "Partial editing"], Inches(0.9), Inches(2.0), Inches(4.2), Inches(2.4), 20)
    quote(s, "How can we detect and explain suspicious forensic evidence in audio instead of only giving a binary fake/real label?", Inches(5.0), Inches(2.0), Inches(7.2), Inches(2.4), CYAN)
    card(s, Inches(1.0), Inches(5.25), Inches(3.5), Inches(0.9), "Risk", "Overclaiming can create false accusations", RED, 15)
    card(s, Inches(4.9), Inches(5.25), Inches(3.5), Inches(0.9), "Need", "Evidence bands + explanation", AMBER, 15)
    card(s, Inches(8.8), Inches(5.25), Inches(3.5), Inches(0.9), "Output", "Manual-review support", GREEN, 15)
    n += 1

    # 4
    s = new_slide(prs, n, "Project Aim", "Move from a label to explainable forensic evidence axes")
    flow(s, ["AI-origin evidence", "Replay / rerecording", "Mixer / channel", "Partial fabrication"], Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.35), CYAN)
    card(s, Inches(1.1), Inches(4.35), Inches(4.8), Inches(1.1), "From", "“Is this fake or real?”", RED, 20)
    card(s, Inches(7.0), Inches(4.35), Inches(4.8), Inches(1.1), "To", "“What evidence exists?”", GREEN, 20)
    n += 1

    # 5
    s = new_slide(prs, n, "Methodology", "Iterative and incremental development driven by experimental failures")
    flow(s, ["Build", "Evaluate", "Find failure", "Improve", "Test again", "Freeze release"], Inches(0.65), Inches(2.1), Inches(12.0), Inches(1.15), AMBER)
    bullet_list(s, ["Requirements changed after real-world testing.", "Each phase produced a working output.", "Final architecture was selected using evidence, not only benchmark scores."], Inches(1.0), Inches(4.3), Inches(10.7), Inches(1.5), 18)
    n += 1

    # 6
    s = new_slide(prs, n, "Full Project Journey", "Every major design change came from a discovered limitation")
    steps = ["CNN/LCNN", "ResNet", "Environmental", "Hybrid", "AASIST", "Controlled tests", "Multi-axis", "Phase 9 API", "Web deploy", "Case studies"]
    flow(s, steps[:5], Inches(0.6), Inches(1.95), Inches(12.1), Inches(1.05), CYAN)
    flow(s, steps[5:], Inches(0.6), Inches(3.45), Inches(12.1), Inches(1.05), GREEN)
    quote(s, "The final system is the result of failure analysis, not random model switching.", Inches(1.1), Inches(5.45), Inches(11.0), Inches(0.85), AMBER)
    n += 1

    # 7 Dataset
    s = new_slide(prs, n, "Dataset Overview", "Large unified dataset with spoof, bona fide, replay, conversion and synthesis samples")
    card(s, Inches(0.8), Inches(1.75), Inches(3.0), Inches(0.9), "Unified size", "1,893,919", CYAN, 24)
    card(s, Inches(4.0), Inches(1.75), Inches(3.0), Inches(0.9), "Spoof", "1,573,308", RED, 24)
    card(s, Inches(7.2), Inches(1.75), Inches(3.0), Inches(0.9), "Bona fide", "320,611", GREEN, 24)
    add_doughnut_chart(s, Inches(0.85), Inches(3.05), Inches(5.3), Inches(3.3), "Spoof vs Bona Fide Distribution", ["Spoof", "Bona fide"], [1573308, 320611])
    add_bar_chart(s, Inches(6.55), Inches(3.05), Inches(5.8), Inches(3.3), "Spoof Type Counts", ["Replay", "Conversion", "Synthesis"], [("Samples (k)", [816.48, 589.212, 167.616])], "Thousands")
    n += 1

    # 8 Pipeline
    s = new_slide(prs, n, "Audio Processing Pipeline", "Consistent backend processing before evidence scoring")
    flow(s, ["Upload", "Load", "Resample", "Normalize", "Segment", "Features", "Evidence models", "Fusion", "Report"], Inches(0.55), Inches(2.15), Inches(12.2), Inches(1.1), CYAN)
    bullet_list(s, ["Standardizes different audio formats and sample rates.", "Segment-level analysis supports partial fabrication review.", "Reports are generated as decision-support evidence, not proof."], Inches(1.0), Inches(4.45), Inches(11), Inches(1.4), 18)
    n += 1

    # 9 Features
    s = new_slide(prs, n, "Feature Extraction Strategy", "Different manipulations leave different traces")
    table_data = [
        ["Feature type", "Primary role"],
        ["LFCC", "Baseline spoofing detection"],
        ["Log-Mel", "Spectral deep learning models"],
        ["Environmental features", "Replay and channel evidence"],
        ["SSL embeddings", "AI-origin evidence"],
        ["Segment transitions", "Partial fabrication candidates"],
    ]
    add_table(s, Inches(1.0), Inches(1.8), Inches(6.0), Inches(3.2), len(table_data), 2, table_data)
    bullet_list(s, ["AI speech: voice-origin patterns", "Replay: recording-chain patterns", "Mixer/channel: acoustic conditions", "Partial fabrication: local segment discontinuities"], Inches(7.6), Inches(2.0), Inches(4.7), Inches(2.8), 18)
    n += 1

    # 10 model comparison
    s = new_slide(prs, n, "Model Experiments", "Strong benchmark numbers did not always mean safe deployment")
    add_bar_chart(s, Inches(0.8), Inches(1.7), Inches(5.6), Inches(3.4), "EER Comparison (lower is better)", ["CNN clean", "CNN aug", "ResNet clean", "ResNet aug", "Hybrid"], [("EER %", [9.68, 15.71, 0.57, 2.61, 16.21])], "%", 100)
    card(s, Inches(7.0), Inches(1.9), Inches(5.2), Inches(0.95), "Key lesson", "ResNet had excellent EER but failed real-world style tests.", AMBER, 16)
    card(s, Inches(7.0), Inches(3.1), Inches(5.2), Inches(0.95), "Decision", "Final system prioritizes evidence coverage and safe reporting.", GREEN, 16)
    card(s, Inches(7.0), Inches(4.3), Inches(5.2), Inches(0.95), "Defense point", "We do not claim universal benchmark superiority.", CYAN, 16)
    n += 1

    # 11 AASIST
    s = new_slide(prs, n, "AASIST Experiment", "A strong literature model was rejected after controlled testing")
    add_bar_chart(s, Inches(1.0), Inches(1.9), Inches(5.3), Inches(3.5), "Clean Human Test Result", ["False flagged", "Not flagged"], [("Count", [22, 1])], "Recordings")
    quote(s, "22 out of 23 clean human recordings were falsely flagged.", Inches(6.9), Inches(2.05), Inches(5.1), Inches(1.2), RED)
    bullet_list(s, ["False accusations are high-risk in forensic interfaces.", "A strong research model is not automatically safe for deployment.", "This justified the cautious multi-axis design."], Inches(7.0), Inches(3.65), Inches(5.0), Inches(1.4), 16)
    n += 1

    # 12 Final architecture
    s = new_slide(prs, n, "Final Phase 9 Architecture", "Four evidence axes + fusion + report layer")
    flow(s, ["Audio", "Preprocess", "Origin", "Replay", "Channel", "Partial", "Fusion", "Report"], Inches(0.55), Inches(2.0), Inches(12.2), Inches(1.12), CYAN)
    card(s, Inches(0.9), Inches(4.25), Inches(2.8), Inches(1.0), "Origin", "Synthetic voice indicators", CYAN, 14)
    card(s, Inches(3.95), Inches(4.25), Inches(2.8), Inches(1.0), "Replay", "Rerecording evidence", AMBER, 14)
    card(s, Inches(7.0), Inches(4.25), Inches(2.8), Inches(1.0), "Channel", "Processing artifacts", PURPLE, 14)
    card(s, Inches(10.05), Inches(4.25), Inches(2.8), Inches(1.0), "Partial", "Suspicious moments", GREEN, 14)
    n += 1

    # 13 axis metrics
    s = new_slide(prs, n, "Final Multi-Axis Release Results", "Separate metrics prevent one weak axis from hiding inside one score")
    add_bar_chart(
        s,
        Inches(0.7),
        Inches(1.65),
        Inches(12.0),
        Inches(4.25),
        "Final Release Matrix Metrics",
        ["Origin", "Replay", "Mixer/channel", "Partial"],
        [
            ("Accuracy", [83.33, 80.00, 88.00, 100.00]),
            ("Balanced accuracy", [82.50, 77.38, 47.83, 100.00]),
            ("Recall", [90.00, 71.43, 0.00, 100.00]),
        ],
        "%",
        100,
    )
    textbox(s, Inches(0.9), Inches(6.15), Inches(11.6), Inches(0.35), "Honest interpretation: origin/replay useful, mixer/channel weak recall, partial promising but still experimental.", 14, True, AMBER, PP_ALIGN.CENTER)
    n += 1

    # 14 Axis details
    s = new_slide(prs, n, "Evidence Axis Logic", "What each axis means in presentation Q&A")
    data = [
        ["Axis", "Threshold", "Meaning", "Important caveat"],
        ["Origin", "0.92", "AI-origin indicators", "Compression can mask cues"],
        ["Replay", "0.65", "Rerecording evidence", "Does not mean AI"],
        ["Mixer/channel", "0.75", "Channel artifacts", "Weak positive recall"],
        ["Partial", "0.95", "Candidate moments", "Manual-review only"],
    ]
    add_table(s, Inches(0.8), Inches(1.75), Inches(11.8), Inches(3.2), len(data), 4, data)
    quote(s, "The system reports evidence indicators, not legal proof.", Inches(1.6), Inches(5.45), Inches(10.2), Inches(0.8), CYAN)
    n += 1

    # 15 Fusion/report
    s = new_slide(prs, n, "Fusion and Report Logic", "Designed to avoid unsafe fake/real overclaiming")
    bullet_list(s, ["Low / moderate / high evidence bands", "Manual review recommendation", "No conclusive authenticity certificate", "Segment timestamps for listening", "Plain-language report + technical export"], Inches(1.0), Inches(1.9), Inches(5.6), Inches(3.4), 18)
    quote(s, "FASSD does not replace forensic experts; it helps them know where to look.", Inches(6.75), Inches(2.15), Inches(5.5), Inches(2.3), GREEN)
    n += 1

    # 16 Literature positioning
    s = new_slide(prs, n, "Literature Positioning", "Our contribution is practical forensic reporting")
    data = [
        ["Common literature focus", "FASSD extension"],
        ["Single spoof score", "Four evidence axes"],
        ["Benchmark-only testing", "Controlled + real-world case testing"],
        ["Model accuracy", "Evidence explanation"],
        ["File-level classification", "Segment candidates"],
        ["Research output", "Deployed web application"],
    ]
    add_table(s, Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.65), len(data), 2, data)
    textbox(s, Inches(1.2), Inches(5.85), Inches(10.9), Inches(0.35), "Defense answer: better in evidence coverage and practical reporting, not universal benchmark superiority.", 15, True, AMBER, PP_ALIGN.CENTER)
    n += 1

    # 17 Deployment
    s = new_slide(prs, n, "Public Web Deployment", "A complete software product, not only notebooks")
    flow(s, ["Next.js / Vercel", "Firebase Auth", "FastAPI API", "Phase 9 Models", "PDF/JSON Reports"], Inches(0.8), Inches(1.9), Inches(11.7), Inches(1.2), CYAN)
    card(s, Inches(1.0), Inches(3.9), Inches(3.5), Inches(1.0), "Frontend", "Dashboard, recorder, test clips, AI lab", BLUE, 14)
    card(s, Inches(4.9), Inches(3.9), Inches(3.5), Inches(1.0), "Backend", "Model inference + report generation", GREEN, 14)
    card(s, Inches(8.8), Inches(3.9), Inches(3.5), Inches(1.0), "Security", "Auth-gated analysis + email verification", AMBER, 14)
    textbox(s, Inches(1.0), Inches(5.65), Inches(11.5), Inches(0.3), "Live: https://www.deepfakedetection.dev", 17, True, CYAN, PP_ALIGN.CENTER)
    n += 1

    # 18 Web user flow
    s = new_slide(prs, n, "Final Interface Output", "What the reviewer receives")
    flow(s, ["Upload / record", "Analyze", "4 evidence cards", "Review segments", "Safety wording", "Export report"], Inches(0.65), Inches(2.1), Inches(12.0), Inches(1.15), GREEN)
    bullet_list(s, ["Readable headline", "Evidence band for each axis", "Highlighted timestamps", "Report-oriented explanation", "Manual-review recommendation"], Inches(1.0), Inches(4.3), Inches(10.8), Inches(1.4), 18)
    n += 1

    # 19 Case studies overview
    s = new_slide(prs, n, "Real-World Case Studies", "Public incidents tested beyond controlled benchmark datasets")
    card(s, Inches(1.0), Inches(2.0), Inches(5.0), Inches(1.2), "Case 1", "Biden New Hampshire robocall", AMBER, 20)
    card(s, Inches(7.2), Inches(2.0), Inches(5.0), Inches(1.2), "Case 2", "Baltimore / Pikesville principal audio", CYAN, 18)
    bullet_list(s, ["Purpose: compare public forensic discussion with FASSD evidence-axis output.", "Safety: experimental evidence only, not legal proof."], Inches(1.25), Inches(4.25), Inches(10.8), Inches(1.25), 18)
    n += 1

    # 20 Biden
    s = new_slide(prs, n, "Case Study 1: Biden Robocall", "A difficult phone-channel/public-media copy")
    add_bar_chart(s, Inches(0.8), Inches(1.75), Inches(6.0), Inches(3.6), "FASSD Evidence Estimates", ["Origin", "Replay", "Channel", "Partial"], [("Estimate %", [88.7, 96.0, 96.0, 96.0])], "%", 100)
    card(s, Inches(7.2), Inches(1.95), Inches(4.8), Inches(0.9), "Agreement", "Mixed / partial", AMBER, 20)
    bullet_list(s, ["External finding: AI-generated robocall.", "FASSD: borderline origin, high replay/channel.", "Interpretation: public copy likely dominated by phone/media artifacts."], Inches(7.25), Inches(3.25), Inches(4.8), Inches(1.7), 15)
    n += 1

    # 21 Pikesville
    s = new_slide(prs, n, "Case Study 2: Pikesville Principal Audio", "A stronger match with public forensic discussion")
    add_bar_chart(s, Inches(0.8), Inches(1.75), Inches(6.0), Inches(3.6), "FASSD Evidence Estimates", ["Origin", "Replay", "Channel", "Partial"], [("Estimate %", [96.0, 0.0, 0.0, 96.0])], "%", 100)
    card(s, Inches(7.2), Inches(1.95), Inches(4.8), Inches(0.9), "Agreement", "Broad agreement", GREEN, 20)
    bullet_list(s, ["External discussion: AI-generated or manipulated.", "FASSD: elevated origin + partial candidates.", "Replay/channel low for this public copy."], Inches(7.25), Inches(3.25), Inches(4.8), Inches(1.7), 15)
    n += 1

    # 22 Case comparison
    s = new_slide(prs, n, "Case Study Comparison", "Multi-axis reporting explains why results differ")
    data = [
        ["Case", "External finding", "FASSD output", "Agreement"],
        ["Biden NH Robocall", "AI-generated robocall", "Borderline origin + high channel", "Mixed"],
        ["Pikesville Principal", "AI-generated / manipulated", "High origin + partial candidates", "Broad"],
    ]
    add_table(s, Inches(0.7), Inches(1.8), Inches(12.0), Inches(2.0), len(data), 4, data)
    bullet_list(s, ["Public copies can lose or distort forensic traces.", "Evidence axes explain inconclusive outputs instead of hiding them.", "Case studies make the demo stronger than benchmark-only evaluation."], Inches(1.0), Inches(4.4), Inches(10.8), Inches(1.3), 18)
    n += 1

    # 23 Firestore
    s = new_slide(prs, n, "Why Firestore NoSQL?", "Simple user history, not heavy ML storage")
    data = [
        ["Reason", "Explanation"],
        ["Firebase Auth integration", "Direct connection with user accounts"],
        ["Flexible schema", "Report fields changed during development"],
        ["Simple history", "Each analysis stored as one document"],
        ["Fast deployment", "No separate SQL server needed"],
        ["ML separation", "FastAPI handles inference"],
    ]
    add_table(s, Inches(1.0), Inches(1.75), Inches(11.3), Inches(3.5), len(data), 2, data)
    textbox(s, Inches(1.0), Inches(5.75), Inches(11.3), Inches(0.35), "Clarification: Firestore stores web history, not the full ML inference pipeline.", 15, True, AMBER, PP_ALIGN.CENTER)
    n += 1

    # 24 Limitations
    s = new_slide(prs, n, "Limitations", "Honesty strengthens the defense")
    bullet_list(s, ["Not court-ready", "Manual review required", "Mixer/channel recall is weak", "WhatsApp/phone compression can hide origin cues", "Partial fabrication is experimental", "Public copies may lose forensic traces"], Inches(1.0), Inches(1.8), Inches(6.2), Inches(3.7), 18)
    quote(s, "FASSD is a research prototype and decision-support tool, not a final forensic authority.", Inches(7.3), Inches(2.2), Inches(4.9), Inches(2.2), RED)
    n += 1

    # 25 Future
    s = new_slide(prs, n, "Future Work", "Path from FYP prototype to stronger forensic platform")
    bullet_list(s, ["Larger real-world forensic dataset", "Better WhatsApp and phone-channel robustness", "Improved mixer/channel model", "Stronger partial localization", "Expert forensic validation", "Chain-of-custody case management"], Inches(1.0), Inches(1.8), Inches(10.8), Inches(3.2), 19)
    quote(s, "First improvement priority: mixer/channel axis and compression robustness.", Inches(1.5), Inches(5.5), Inches(10.2), Inches(0.8), AMBER)
    n += 1

    # 26 Contributions
    s = new_slide(prs, n, "Final Contributions", "What FASSD adds as a complete FYP")
    bullet_list(s, ["Complete audio deepfake detection pipeline", "Large unified dataset preparation", "Multiple model experiments and documented failures", "Controlled forensic testing", "Final multi-axis evidence architecture", "Evidence-band reporting system", "Public deployed web application", "Real-world case study evaluation"], Inches(0.9), Inches(1.65), Inches(6.4), Inches(4.3), 16)
    quote(s, "Biggest contribution: moving from binary detection to multi-axis forensic evidence reporting.", Inches(7.45), Inches(2.15), Inches(4.8), Inches(2.4), GREEN)
    n += 1

    # 27 Conclusion
    s = new_slide(prs, n, "Conclusion", "A complete ML + software engineering project")
    bullet_list(s, ["Achieved official FYP objectives", "Extended into deployed forensic decision-support", "Analyzes four evidence axes", "Provides readable reports and segment highlights", "Avoids unsafe fake/real overclaiming"], Inches(1.0), Inches(1.8), Inches(10.9), Inches(2.8), 19)
    quote(s, "FASSD does not replace forensic experts; it helps them know where to look.", Inches(1.4), Inches(5.15), Inches(10.5), Inches(0.9), CYAN)
    n += 1

    # 28 Thank you
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    textbox(s, Inches(0.9), Inches(1.3), Inches(11.5), Inches(0.8), "Thank You", 50, True, TEXT, PP_ALIGN.CENTER)
    textbox(s, Inches(0.9), Inches(2.25), Inches(11.5), Inches(0.45), "Questions?", 28, True, CYAN, PP_ALIGN.CENTER)
    card(s, Inches(2.0), Inches(4.0), Inches(4.5), Inches(0.9), "Live website", "https://www.deepfakedetection.dev", CYAN, 14)
    card(s, Inches(6.9), Inches(4.0), Inches(4.5), Inches(0.9), "Repositories", "github.com/RanaAreeb/FASSD\ngithub.com/MHS-4ever/FYP_FASSD", GREEN, 12)
    footer(s, n)
    n += 1

    # Appendix Q&A
    s = new_slide(prs, n, "Backup Viva Questions", "High-risk questions and compact answers")
    data = [
        ["Question", "Defense answer"],
        ["Why not one model?", "Different manipulations leave different traces."],
        ["Why reject ResNet?", "Great benchmark EER, weak real-world generalization."],
        ["Why reject AASIST?", "22/23 clean human false alarms in controlled test."],
        ["Can FASSD prove fake?", "No. It provides decision-support evidence only."],
        ["Biggest limitation?", "Mixer/channel recall and compression robustness."],
    ]
    add_table(s, Inches(0.65), Inches(1.65), Inches(12.0), Inches(4.4), len(data), 2, data)
    n += 1

    s = new_slide(prs, n, "Backup Viva Questions", "More defense-ready answers")
    data = [
        ["Question", "Defense answer"],
        ["Replay vs AI-origin?", "Replay is recording-chain; AI-origin is voice source."],
        ["Why evidence bands?", "Safer and clearer than raw confidence scores."],
        ["Better than literature?", "Stronger practical reporting; no universal superiority claim."],
        ["Why mixed Biden result?", "Phone/media channel artifacts masked origin cues."],
        ["Why manual review?", "Audio forensics is high-risk; automated models can fail."],
    ]
    add_table(s, Inches(0.65), Inches(1.65), Inches(12.0), Inches(4.4), len(data), 2, data)

    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
